import asyncio
import base64
import csv
import io
import json
import math
import re
import uuid
import sys
from html import escape
from html.parser import HTMLParser
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, Reference
import httpx
from .config import settings
from .db import File, Session
from sqlalchemy import select
from .schemas import TOOLS
from .security import fetch_public, public_address
from .storage import storage


class TextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.hidden += 1

    def handle_endtag(self, tag):
        if tag in {"script", "style"}:
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden and data.strip():
            self.parts.append(data.strip())


def task_file(task_id, file_id):
    with Session() as db:
        f = db.get(File, file_id)
        if not f or f.task_id != task_id:
            raise ValueError("File does not belong to this task")
        return f, storage.read(f.storage_key)


def csv_rows(content):
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames or len(set(reader.fieldnames)) != len(reader.fieldnames):
        raise ValueError("CSV requires unique column headers")
    if len(reader.fieldnames) > 100:
        raise ValueError("CSV exceeds 100 columns")
    rows = []
    for row in reader:
        if len(rows) >= 100000 or None in row:
            raise ValueError("CSV exceeds row limit or contains malformed rows")
        rows.append(row)
    return reader.fieldnames, rows


def ingest(task_id, args):
    f, content = task_file(task_id, args.file_id)
    if f.name.lower().endswith(".csv"):
        columns, rows = csv_rows(content)
        return {"columns": columns, "rows": len(rows), "sample": rows[:5], "origin": "uploaded data"}, []
    if f.name.lower().endswith(".pdf"):
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted or len(reader.pages) > 100:
            raise ValueError("PDF must be unencrypted and at most 100 pages")
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        if not text.strip():
            raise ValueError("PDF has no extractable text; OCR is not supported")
    else:
        text = content.decode("utf-8-sig")
    return {"text": text[:16000], "truncated": len(text) > 16000, "origin": "uploaded data"}, []


def artifact(name, mime, content):
    return {"id": str(uuid.uuid4()), "name": name, "mime": mime, "size": len(content), "storage_key": storage.put(content), "kind": "artifact"}


def safe_cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def analyze(task_id, args):
    _, content = task_file(task_id, args.file_id)
    columns, rows = csv_rows(content)
    if args.category_column not in columns or args.value_column not in columns:
        raise ValueError(f"Required columns missing: {args.category_column}, {args.value_column}; available: {columns}")
    totals = {}
    excluded = 0
    for row in rows:
        category = (row.get(args.category_column) or "").strip()
        try:
            value = float(row[args.value_column])
        except (ValueError, TypeError):
            excluded += 1
            continue
        if not category or not math.isfinite(value):
            excluded += 1
            continue
        if len(category) > 120:
            raise ValueError("Category labels must be at most 120 characters")
        totals[category] = totals.get(category, 0) + value
    if not totals or len(totals) > 200:
        raise ValueError("Analysis requires 1–200 categories with finite numeric values")
    if not all(math.isfinite(value) for value in totals.values()):
        raise ValueError("Category totals overflowed; scale the input values")
    ranked = [{"category": k, "value": round(v, 4)} for k, v in sorted(totals.items(), key=lambda kv: (kv[1], kv[0]))]
    with Session() as db:
        csv_count = sum(f.name.lower().endswith(".csv") for f in db.scalars(select(File).where(File.task_id == task_id, File.kind == "upload")))
    suffix = "-" + args.file_id[:8] if csv_count > 1 else ""
    chart_name = "category-revenue" + suffix + ".svg"
    shown = ranked[:20]
    max_abs = max(abs(r["value"]) for r in shown) or 1
    height = 100 + 36 * len(shown)
    bars = []
    for i, row in enumerate(shown):
        y = 70 + i * 36
        width = abs(row["value"]) / max_abs * 180
        x = 420 - width if row["value"] < 0 else 420
        bars.append(
            f'<text x="24" y="{y + 17}" fill="#414a39">{escape(row["category"][:26])}</text><rect x="{x}" y="{y}" width="{width}" height="24" rx="2" fill="{"#a96543" if i < 3 else "#c1cab7"}"/><text x="620" y="{y + 17}" fill="#414a39">{row["value"]:,.2f}</text>'
        )
    svg = f'<svg xmlns="http://www.w3.org/2000/svg" width="800" height="{height}" viewBox="0 0 800 {height}"><rect width="800" height="{height}" fill="#ffffff"/><g font-family="Arial, sans-serif" font-size="14"><text x="24" y="30" font-size="18" fill="#343d2c">Total {escape(args.value_column)} by {escape(args.category_column)} (lowest {len(shown)})</text><line x1="420" y1="55" x2="420" y2="{height - 20}" stroke="#e1e5da"/>{"".join(bars)}</g></svg>'
    book = Workbook()
    sheet = book.active
    sheet.title = "Category totals"
    sheet.append([f"Source: uploaded file {args.file_id}", None])
    sheet.append([args.category_column, f"Sum of {args.value_column}"])
    for row in ranked:
        sheet.append([safe_cell(row["category"]), row["value"]])
        sheet.cell(sheet.max_row, 2).number_format = "#,##0.00"
    sheet.append(["Excluded rows", excluded])
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:B{len(ranked) + 2}"
    sheet.column_dimensions["A"].width = 48
    sheet.column_dimensions["B"].width = 24
    for cell in sheet[2]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="3345A8")
    chart = BarChart()
    chart.title = "Category revenue"
    chart.add_data(Reference(sheet, min_col=2, min_row=2, max_row=len(ranked) + 2), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=3, max_row=len(ranked) + 2))
    sheet.add_chart(chart, "D2")
    out = io.BytesIO()
    book.save(out)
    return {
        "metric": f"Ascending sum of {args.value_column} by {args.category_column}; ties alphabetical. Negative values retained. Not profit or growth.",
        "rows": len(rows),
        "excluded_rows": excluded,
        "weakest": ranked[:3],
        "totals": ranked,
        "origin": "uploaded data",
        "file_id": args.file_id,
        "chart_name": chart_name,
    }, [
        artifact(chart_name, "image/svg+xml", svg.encode()),
        artifact("category-analysis" + suffix + ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", out.getvalue()),
    ]


async def search(args, mode):
    if mode == "demo":
        return {
            "sources": [
                {
                    "id": "fixture-market",
                    "title": "Sample market context · DEMO FIXTURE",
                    "url": "https://example.com",
                    "excerpt": "Fictional sample: category demand may differ by season. This is not evidence about any real market.",
                    "fixture": True,
                }
            ],
            "origin": "fixture, not live research",
        }
    cfg = settings()
    if cfg.search_provider == "brave" and not cfg.search_api_key:
        raise ValueError("SEARCH_API_KEY is required for live research (Brave Search)")
    async with httpx.AsyncClient(timeout=20) as client:
        if cfg.search_provider == "tavily":
            headers = {"Authorization": f"Bearer {cfg.tavily_api_key}"} if cfg.tavily_api_key else {"X-Tavily-Access-Mode": "keyless"}
            response = await client.post(
                "https://api.tavily.com/search",
                headers=headers,
                json={"query": args.query, "max_results": 5, "search_depth": "basic", "auto_parameters": False, "include_answer": False},
            )
        else:
            response = await client.get(
                "https://api.search.brave.com/res/v1/web/search", params={"q": args.query, "count": 5}, headers={"X-Subscription-Token": cfg.search_api_key}
            )
        if response.status_code in {429, 432, 433}:
            raise ValueError("Search quota or rate limit reached; retry later or configure a search API key with available quota")
        response.raise_for_status()
        payload = response.json()
        results = (payload.get("results", []) if cfg.search_provider == "tavily" else payload.get("web", {}).get("results", []))[:5]
    sources = []
    for result in results:
        try:
            page = await asyncio.to_thread(fetch_public, result["url"])
            parser = TextExtractor()
            parser.feed(page["body"])
            sources.append(
                {"id": str(uuid.uuid4()), "title": result["title"][:200], "url": page["url"], "excerpt": " ".join(parser.parts)[:2500], "fixture": False}
            )
        except (ValueError, OSError):
            continue
    if not sources:
        raise ValueError("No retrievable public sources; try a narrower search")
    return {"sources": sources, "origin": "public web excerpts; claims require support"}


async def browse(args, mode):
    if mode == "demo":
        return {
            "sources": [
                {
                    "id": "fixture-browser",
                    "title": "Example page · DEMO FIXTURE",
                    "url": "https://example.com",
                    "excerpt": "Example Domain. This domain is for illustrative examples.",
                    "fixture": True,
                }
            ],
            "origin": "fixture; no browser was launched",
        }
    await asyncio.to_thread(public_address, args.url)
    from playwright.async_api import async_playwright

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True, args=["--disable-dev-shm-usage"])
        try:
            context = await browser.new_context(java_script_enabled=False, service_workers="block", accept_downloads=False)
            final_url = args.url

            async def route_request(route):
                nonlocal final_url
                request = route.request
                if request.method != "GET" or request.resource_type != "document":
                    await route.abort()
                    return
                try:
                    page = await asyncio.to_thread(fetch_public, request.url)
                    final_url = page["url"]
                    await route.fulfill(
                        status=200, body=page["body"], content_type=page["mime"], headers={"Content-Security-Policy": "default-src 'none'; script-src 'none'"}
                    )
                except (ValueError, OSError):
                    await route.abort()

            await context.route("**/*", route_request)
            page = await context.new_page()
            await page.goto(args.url, wait_until="domcontentloaded", timeout=30000)
            text = (await page.locator("body").inner_text())[:10000]
            title = (await page.title())[:200]
            return {
                "sources": [{"id": str(uuid.uuid4()), "title": title, "url": final_url, "excerpt": text, "fixture": False}],
                "origin": "public browser; scripts disabled",
            }
        finally:
            await browser.close()


def report(args, checkpoint, mode):
    known = {s["id"]: s for r in checkpoint.get("context_results", []) + checkpoint.get("results", []) for s in r["output"].get("sources", [])}
    if any(s not in known for s in args.source_ids):
        raise ValueError("Report contains unknown source IDs")
    sources = [known[s] for s in args.source_ids]
    label = "DEMO · fixture-backed workflow" if mode == "demo" else "LIVE · model-generated synthesis; verify interpretations"
    markdown = f"# {args.title}\n\n{label}\n\n## Findings\n{args.findings}\n\n## Hypotheses and suggested experiments\n{args.hypotheses}\n"
    styles = getSampleStyleSheet()
    story = [Paragraph(escape(args.title), styles["Title"]), Paragraph(escape(label), styles["Normal"]), Spacer(1, 18)]
    for title, text in [("Findings from data / observations", args.findings), ("Hypotheses and suggested experiments", args.hypotheses)]:
        story.append(Paragraph(title, styles["Heading2"]))
        for line in text.splitlines():
            story.append(Paragraph(escape(line), styles["BodyText"]))
    for r in checkpoint.get("results", []):
        if r["tool"] == "analyze_csv" and "totals" in r["output"]:
            output = r["output"]
            markdown += (
                "\n## Calculation method\n"
                + output["metric"]
                + f" Excluded rows: {output['excluded_rows']}. Input file: {output.get('file_id', '')}.\n\n![Category revenue]({output.get('chart_name', 'category-revenue.svg')})\n"
            )
            story.extend([Paragraph("Category performance", styles["Heading2"]), Paragraph(escape(output["metric"]), styles["BodyText"])])
            # Vector bars embedded in the PDF; signed totals and baseline retained.
            from reportlab.graphics.shapes import Drawing, Rect, String, Line

            shown = output["totals"][:15]
            scale = max(abs(row["value"]) for row in shown) or 1
            drawing = Drawing(460, 30 * len(shown) + 15)
            drawing.add(Line(260, 0, 260, drawing.height, strokeColor=colors.lightgrey))
            for i, row in enumerate(shown):
                y = drawing.height - 25 - i * 30
                width = abs(row["value"]) / scale * 100
                drawing.add(String(0, y + 5, row["category"][:25], fontSize=9))
                drawing.add(Rect(260 - width if row["value"] < 0 else 260, y, width, 18, fillColor=colors.HexColor("#a96543" if i < 3 else "#c1cab7"), strokeColor=None))
                drawing.add(String(370, y + 5, f"{row['value']:,.2f}", fontSize=9))
            story.append(drawing)
            table = Table(
                [["Category", "Total"]] + [[Paragraph(escape(row["category"]), styles["Normal"]), f"{row['value']:,.2f}"] for row in output["weakest"]],
                colWidths=[300, 140],
            )
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#edf1e7")), ("BOTTOMPADDING", (0, 0), (-1, -1), 10)]))
            story.append(table)
    markdown += "\n## External source excerpts\n"
    for source in sources:
        markdown += f"\n[{source['id']}] {source['title']}\n{source['url']}\n\n> {source['excerpt']}\n"
        story.extend(
            [
                Paragraph(escape(source["title"]), styles["Heading2"]),
                Paragraph(escape(source["url"]), styles["BodyText"]),
                Paragraph(escape(source["excerpt"]), styles["BodyText"]),
            ]
        )
    if not sources:
        markdown += "\nNo external sources were used.\n"
    out = io.BytesIO()

    def footer(canvas, doc):
        canvas.setFont("Helvetica", 9)
        canvas.drawString(48, 28, f"AutoAgent | {mode.upper()} | Page {doc.page}")

    SimpleDocTemplate(out, rightMargin=48, leftMargin=48, topMargin=42, bottomMargin=48).build(story, onFirstPage=footer, onLaterPages=footer)
    return {"summary": args.title, "source_ids": args.source_ids}, [
        artifact("report.md", "text/markdown", markdown.encode()),
        artifact("report.pdf", "application/pdf", out.getvalue()),
    ]


async def python_tool(task_id, args):
    files = {f: base64.b64encode(task_file(task_id, f)[1]).decode() for f in args.file_ids}
    async with httpx.AsyncClient(timeout=65) as client:
        response = await client.post(
            settings().sandbox_url + "/execute", headers={"Authorization": "Bearer " + settings().sandbox_secret}, json={"code": args.code, "files": files}
        )
        response.raise_for_status()
        result = response.json()
    if result.get("error"):
        raise ValueError(result["error"])
    allowed = {".png": "image/png", ".csv": "text/csv", ".txt": "text/plain", ".json": "application/json"}
    artifacts = []
    for file in result.get("files", []):
        ext = "." + file["name"].rsplit(".", 1)[-1].lower()
        if ext not in allowed or not re.fullmatch(r"[a-zA-Z0-9_.-]{1,100}", file["name"]):
            continue
        content = base64.b64decode(file["content"], validate=True)
        if len(content) > 5_000_000:
            raise ValueError("Sandbox artifact exceeds size limit")
        artifacts.append(artifact(file["name"], allowed[ext], content))
    return {"stdout": result.get("stdout", "")[:12000]}, artifacts


async def execute(task_id, tool, arguments, checkpoint, mode):
    args = TOOLS[tool].model_validate(arguments)
    if tool in {"ingest", "analyze_csv", "report"}:
        proc = await asyncio.create_subprocess_exec(
            sys.executable, "-m", "app.trusted_worker", stdin=asyncio.subprocess.PIPE, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.DEVNULL
        )
        try:
            stdout, _ = await proc.communicate(
                json.dumps({"task_id": task_id, "tool": tool, "arguments": arguments, "checkpoint": checkpoint, "mode": mode}).encode()
            )
            if proc.returncode:
                raise ValueError("Fixed tool process failed or exceeded resource limits")
            decoded = json.loads(stdout)
            if "error" in decoded:
                raise ValueError(decoded["error"])
            result = decoded["result"]
        finally:
            if proc.returncode is None:
                proc.kill()
                await proc.wait()
    elif tool == "search":
        result = (await search(args, mode), [])
    elif tool == "browse":
        result = (await browse(args, mode), [])
    elif tool == "python":
        result = await python_tool(task_id, args)
    else:
        result = ({"summary": args.summary}, [])
    if len(json.dumps(result[0]).encode()) > settings().max_tool_bytes:
        raise ValueError("Tool output exceeds configured byte budget; narrow the request")
    return result
