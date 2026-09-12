"""Editable Word export using the same validated report data as PDF/Markdown."""
import io
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def word_report(args, checkpoint, sources, label):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.8)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    doc.add_paragraph(args.title, "Title")
    doc.add_paragraph(label)

    def paragraphs(text):
        for line in text.splitlines():
            if not line.strip():
                continue
            # Keep citation IDs intact; translate common Markdown paragraph forms.
            heading = re.match(r"^#{1,6}\s+(.+)", line)
            bullet = re.match(r"^\s*[-*+]\s+(.+)", line)
            numbered = re.match(r"^\s*\d+\.\s+(.+)", line)
            value = (heading or bullet or numbered).group(1) if heading or bullet or numbered else line
            style = "Heading 2" if heading else "List Bullet" if bullet else "List Number" if numbered else "Normal"
            paragraph = doc.add_paragraph(style=style)
            for part in re.split(r"(\*\*.*?\*\*)", value):
                run = paragraph.add_run(part[2:-2] if part.startswith("**") and part.endswith("**") else part)
                run.bold = part.startswith("**") and part.endswith("**")

    doc.add_heading("Findings", level=1)
    paragraphs(args.findings)
    doc.add_heading("Hypotheses and suggested experiments", level=1)
    paragraphs(args.hypotheses)
    for result in checkpoint.get("results", []):
        output = result["output"]
        if result["tool"] == "analyze_csv" and "totals" in output:
            doc.add_heading("Category performance", level=1)
            doc.add_paragraph(output["metric"])
            doc.add_paragraph(f"Excluded rows: {output['excluded_rows']}. Input file: {output.get('file_id', '')}.")
            table = doc.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            borders = OxmlElement("w:tblBorders")
            for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
                border = OxmlElement("w:" + edge)
                for name, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                    border.set(qn("w:" + name), value)
                borders.append(border)
            table._tbl.tblPr.append(borders)
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            table.rows[0].cells[0].text = "Category"
            table.rows[0].cells[1].text = "Total"
            for cell in table.rows[0].cells:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EEEEEE")
                cell._tc.get_or_add_tcPr().append(shading)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in output["totals"]:
                cells = table.add_row().cells
                cells[0].text = row["category"]
                cells[1].text = f"{row['value']:,.2f}"
    doc.add_heading("External source excerpts", level=1)
    for source in sources:
        doc.add_heading(source["title"], level=2)
        doc.add_paragraph(f"[{source['id']}]")
        doc.add_paragraph(source["url"])
        doc.add_paragraph(source["excerpt"])
    if not sources:
        doc.add_paragraph("No external sources were used.")
    section.footer.paragraphs[0].text = "AutoQA | " + label
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
