import asyncio
import io
from pypdf import PdfReader
from openpyxl import load_workbook
from app.runtime import run_task

CSV = b"category,revenue\nBooks,20\nGames,5\nFood,12\nBooks,3\nClothes,40\nFood,invalid\n"


def create(client, goal, content=None, filename="sales.csv"):
    ids = []
    if content:
        response = client.post("/api/uploads", files={"file": (filename, content)})
        assert response.status_code == 201, response.text
        ids = [response.json()["id"]]
    response = client.post("/api/tasks", json={"goal": goal, "file_ids": ids})
    assert response.status_code == 201, response.text
    return response.json()["id"]


def test_csv_research_report(client):
    task_id = create(client, "Analyze CSV and research market context; create a report with charts", CSV)
    asyncio.run(run_task(task_id))
    task = client.get(f"/api/tasks/{task_id}").json()
    assert task["status"] == "completed", task
    assert task["usage"]["tokens"] == 0
    assert task["sources"][0]["fixture"] is True
    assert len(task["files"]) == 6
    saved = next(t for t in client.get("/api/tasks").json() if t["id"] == task_id)
    assert saved["artifact_count"] == 5
    for f in task["files"]:
        data = client.get("/api/files/" + f["id"]).content
        if f["name"].endswith(".docx"):
            from docx import Document
            document = Document(io.BytesIO(data))
            word_text = "\n".join(p.text for p in document.paragraphs)
            assert "DEMO" in word_text and "Fictional sample" in word_text
            assert document.tables
        if f["name"].endswith(".pdf"):
            text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)
            assert "DEMO" in text and "Games (5.00)" in text
            assert "Fictional sample" in text
        if f["name"].endswith(".xlsx"):
            book = load_workbook(io.BytesIO(data))
            rows = list(book.active.values)
            assert rows[2] == ("Games", 5)
            assert rows[3] == ("Food", 12)
            assert rows[-1] == ("Excluded rows", 1)
            assert len(book.active._charts) == 1
    stream = client.get(f"/api/tasks/{task_id}/events").text
    assert "category" in stream and "data:" in stream
    last_id = int([line[4:] for line in stream.splitlines() if line.startswith("id: ")][-1])
    assert "data:" not in client.get(f"/api/tasks/{task_id}/events?after={last_id}").text
    # Duplicate deliveries cannot execute a completed task again.
    asyncio.run(run_task(task_id))
    assert client.get(f"/api/tasks/{task_id}").json()["usage"] == task["usage"]


def test_research_report(client):
    task_id = create(client, "Research retail and generate a cited report")
    asyncio.run(run_task(task_id))
    task = client.get(f"/api/tasks/{task_id}").json()
    assert task["status"] == "completed"
    assert len(task["sources"]) == 1
    assert {f["name"] for f in task["files"]} == {"report.pdf", "report.md", "report.docx"}


def test_csv_chart_spreadsheet(client):
    task_id = create(client, "Analyze this CSV and generate a chart and spreadsheet", CSV)
    asyncio.run(run_task(task_id))
    task = client.get(f"/api/tasks/{task_id}").json()
    assert task["status"] == "completed"
    assert task["sources"] == []
    assert any(f["name"].endswith(".svg") for f in task["files"])


def test_malformed_csv_stops_without_success(client):
    task_id = create(client, "Analyze this sales CSV", b"category,revenue\nFood,3,extra\n")
    asyncio.run(run_task(task_id))
    task = client.get(f"/api/tasks/{task_id}").json()
    assert task["status"] == "failed"
    assert "twice" in task["error"]
    assert task["usage"]["steps"] == 2


def test_text_ingestion(client):
    task_id = create(client, "Summarize this file in a report", b"Team notes: launch moved to October.", "notes.txt")
    asyncio.run(run_task(task_id))
    task = client.get(f"/api/tasks/{task_id}").json()
    assert task["status"] == "completed"
    report = next(f for f in task["files"] if f["name"] == "report.md")
    assert "launch moved to October" in client.get("/api/files/" + report["id"]).text


def test_multiple_uploads_are_all_used(client):
    first = client.post("/api/uploads", files={"file": ("first.csv", CSV)}).json()["id"]
    second = client.post("/api/uploads", files={"file": ("second.csv", b"category,revenue\nA,8\nB,2\n")}).json()["id"]
    task_id = client.post("/api/tasks", json={"goal": "Analyze all CSV files and create a report", "file_ids": [first, second]}).json()["id"]
    asyncio.run(run_task(task_id))
    task = client.get("/api/tasks/" + task_id).json()
    assert task["status"] == "completed"
    names = [f["name"] for f in task["files"] if f["kind"] == "artifact"]
    assert len(names) == len(set(names)) == 7
    report = next(f for f in task["files"] if f["name"] == "report.md")
    content = client.get("/api/files/" + report["id"]).text
    assert first in content and second in content
