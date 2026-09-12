import io
from types import SimpleNamespace
from docx import Document
from app.word_report import word_report


def test_word_report_preserves_citations_and_numeric_data():
    args = SimpleNamespace(title="Sales findings", findings="**Revenue** is 42 [s1].", hypotheses="Test positioning.")
    checkpoint = {"results": [{"tool": "analyze_csv", "output": {"totals": [{"category": "Tools", "value": -42}], "metric": "Sum of revenue", "excluded_rows": 2}}]}
    content = word_report(args, checkpoint, [{"id": "s1", "title": "Evidence", "url": "https://example.com", "excerpt": "Revenue is 42."}], "LIVE")
    doc = Document(io.BytesIO(content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Revenue is 42 [s1]." in text
    assert "https://example.com" in text and "Revenue is 42." in text
    assert doc.paragraphs[0].style.name == "Title"
    assert doc.tables[0].rows[1].cells[1].text == "-42.00"
